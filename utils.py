"""
utils.py — Helper utilities for saving data, Excel formatting, and logging.
"""

import os
import logging
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── Logging ──────────────────────────────────────────────────────────────────

def setup_logger(name: str = "jll_scraper", log_file: str = "scraper.log") -> logging.Logger:
    """Configure and return a logger that writes to both console and file."""
    logger = logging.getLogger(name)
    logger.setLevel(logging.DEBUG)

    fmt = logging.Formatter("%(asctime)s [%(levelname)s] %(message)s", "%Y-%m-%d %H:%M:%S")

    # Console handler
    ch = logging.StreamHandler()
    ch.setLevel(logging.INFO)
    ch.setFormatter(fmt)

    # File handler
    fh = logging.FileHandler(log_file, encoding="utf-8")
    fh.setLevel(logging.DEBUG)
    fh.setFormatter(fmt)

    if not logger.handlers:
        logger.addHandler(ch)
        logger.addHandler(fh)

    return logger


# ── Output directory ──────────────────────────────────────────────────────────

def ensure_output_dir(path: str = "output") -> str:
    """Create output directory if it doesn't exist and return its path."""
    os.makedirs(path, exist_ok=True)
    return path


# ── Data cleaning ─────────────────────────────────────────────────────────────

def clean_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """Remove duplicates, strip whitespace, and replace empty strings with NaN."""
    df = df.drop_duplicates(subset=["Property Link"], keep="first")
    df = df.map(lambda x: x.strip() if isinstance(x, str) else x)
    df = df.replace("", pd.NA)
    df.reset_index(drop=True, inplace=True)
    return df


# ── Save helpers ──────────────────────────────────────────────────────────────

def save_csv(df: pd.DataFrame, output_dir: str, city: str, prefix: str = "Combined") -> str:
    """Save DataFrame to CSV and return the file path."""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{prefix}_{city}_offices_{timestamp}.csv"
    path = os.path.join(output_dir, filename)
    df.to_csv(path, index=False, encoding="utf-8-sig")
    return path


def save_excel_multi(sheets: dict, output_dir: str, city: str) -> str:
    """
    Save multiple DataFrames into one Excel workbook.
    sheets = {"Sheet Name": dataframe, ...}
    Returns the file path.
    """
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"Combined_{city}_offices_{timestamp}.xlsx"
    path = os.path.join(output_dir, filename)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, df in sheets.items():
            df.to_excel(writer, index=False, sheet_name=sheet_name[:31])

    _format_excel_multi(path, city, timestamp, sheets)
    return path


def save_excel(df: pd.DataFrame, output_dir: str, city: str) -> str:
    """Save single DataFrame to a formatted Excel file."""
    return save_excel_multi({"Listings": df}, output_dir, city)


def _apply_sheet_format(ws, source_label: str = "") -> None:
    """Apply formatting to a single worksheet."""
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = border
    ws.row_dimensions[1].height = 30

    light_blue = PatternFill("solid", fgColor="D6E4F0")
    white      = PatternFill("solid", fgColor="FFFFFF")
    data_font  = Font(size=10)
    data_align = Alignment(vertical="center", wrap_text=True)

    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        fill = light_blue if row_idx % 2 == 0 else white
        for cell in row:
            cell.fill = fill
            cell.font = data_font
            cell.alignment = data_align
            cell.border = border

    for col_idx, col in enumerate(ws.columns, start=1):
        max_len = max((len(str(c.value)) for c in col if c.value), default=10)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 50)

    ws.freeze_panes = "A2"


def _format_excel_multi(path: str, city: str, timestamp: str, sheets: dict) -> None:
    """Apply formatting to all sheets and add a Metadata sheet."""
    wb = load_workbook(path)

    for sheet_name in sheets:
        ws = wb[sheet_name[:31]]
        _apply_sheet_format(ws)

    # Metadata sheet
    meta = wb.create_sheet("Metadata")
    meta["A1"], meta["B1"] = "Key", "Value"
    meta["A2"], meta["B2"] = "City", city
    meta["A3"], meta["B3"] = "Scraped At", timestamp
    row = 4
    for sheet_name, df in sheets.items():
        meta.cell(row, 1).value = f"{sheet_name} Count"
        meta.cell(row, 2).value = len(df)
        row += 1
    for cell in ["A1", "B1"]:
        meta[cell].font = Font(bold=True)

    wb.save(path)
